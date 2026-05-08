from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]

FILES = [
    ("HANDOUT_SPARRECHNER.md", "HANDOUT_SPARRECHNER.docx", "Sparpotenzial Rechner"),
    ("HANDOUT_TAGESUEBERSICHT.md", "HANDOUT_TAGESUEBERSICHT.docx", "Tagesuebersicht"),
    ("HANDOUT_KETTENVERGLEICH.md", "HANDOUT_KETTENVERGLEICH.docx", "Ketten und freie Tankstellen"),
]


def set_cell_shading(paragraph, color="EEF4F8"):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    p_pr.append(shading)


def set_paragraph_border(paragraph, color="D4DEE7"):
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    border.append(bottom)
    p_pr.append(border)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Seite ")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def apply_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(35, 42, 51)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(21, 89, 126)
    title.paragraph_format.space_after = Pt(8)

    for name, size, before, after in [
        ("Heading 1", 16, 12, 6),
        ("Heading 2", 14, 10, 4),
        ("Heading 3", 12, 8, 3),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(21, 89, 126)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(12)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(6)


def add_run_with_inline_code(paragraph, text):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(72, 85, 99)
        else:
            run = paragraph.add_run(part)
            run.font.name = "Arial"


def render_markdown(document, markdown):
    in_code = False
    code_lines = []

    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()

        if line.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.right_indent = Inches(0.25)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(8)
                set_cell_shading(paragraph)
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Courier New"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(37, 50, 65)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not line.strip():
            continue

        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            add_run_with_inline_code(paragraph, line[2:].strip())
        elif line.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            add_run_with_inline_code(paragraph, line[3:].strip())
        elif line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_run_with_inline_code(paragraph, line[4:].strip())
        elif line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.right_indent = Inches(0.25)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(8)
            set_paragraph_border(paragraph)
            run = paragraph.add_run(line[2:].strip())
            run.italic = True
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(72, 85, 99)
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_run_with_inline_code(paragraph, line[2:].strip())
        elif re.match(r"^\d+\.\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            add_run_with_inline_code(paragraph, re.sub(r"^\d+\.\s+", "", line).strip())
        else:
            paragraph = document.add_paragraph()
            add_run_with_inline_code(paragraph, line)


def build_document(markdown_path, output_path, label):
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    apply_styles(document)

    header = section.header.paragraphs[0]
    header.text = label
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(104, 116, 131)
    set_paragraph_border(header)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(104, 116, 131)

    markdown = markdown_path.read_text(encoding="utf-8")
    render_markdown(document, markdown)
    document.save(output_path)


def main():
    for source, target, label in FILES:
        build_document(ROOT / source, ROOT / target, label)
        print(target)


if __name__ == "__main__":
    main()
